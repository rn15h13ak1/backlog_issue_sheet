"""
Excel の読み書き
================
読み込みは「登録」「更新」の 2 シートと、どのプロジェクト向けかを書く
「プロジェクト」シートを見る。ほかのシートは無視する。

書き込みはひな形・書き出し・実行結果の 3 つで共通に使う。どれも
「更新」シートの書式で出すため、そのまま取り込みに戻せる。
"""

from __future__ import annotations

import datetime as dt
from dataclasses import dataclass, field
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

from fields import (
    CHILD_HEADER,
    CHILD_MARK,
    CLEAR_TOKEN,
    COMMENT_PREFIX,
    DETACH_TOKEN,
    KEY_HEADER,
    PARENT_HEADER,
    REGISTER_SHEET,
    UPDATE_SHEET,
    Column,
    all_columns,
    is_blank,
)
from master import Master

MASTER_SHEET = "_マスタ"
GUIDE_SHEET = "使い方"
PROJECT_SHEET = "プロジェクト"
PROJECT_LABEL = "プロジェクトキー"
PROJECT_NOTE = "取り込むとき、設定ファイルの project_key と照らし合わせます。違えば何も送りません。"
# 入力規則（プルダウン）を付ける行数
VALIDATION_ROWS = 1000


class SheetError(Exception):
    """Excel を読めなかった。"""


@dataclass
class SheetRow:
    row_no: int
    values: dict[int, object]          # 列番号 → セルの値


@dataclass
class SheetData:
    name: str
    headers: list[tuple[int, str]]     # (列番号, 見出し)。# で始まる列と空の見出しは除く
    rows: list[SheetRow] = field(default_factory=list)


class Book(dict):
    """
    {シート名: SheetData}。読み込んだブックに書かれたプロジェクトキーも持つ。

    project_key は「プロジェクト」シートが無ければ None、あっても値が空なら ""。
    """

    def __init__(self, *args, project_key: str | None = None, **kwargs):
        super().__init__(*args, **kwargs)
        self.project_key = project_key


def read_workbook(path: str | Path) -> Book:
    """「登録」「更新」シートとプロジェクトキーを読む。どちらのシートも無ければ SheetError。"""
    path = Path(path)
    if not path.is_file():
        raise SheetError(f"ファイルが見つかりません: {path}")
    if path.suffix.lower() not in (".xlsx", ".xlsm"):
        raise SheetError(f"xlsx / xlsm 以外は読めません: {path.name}")
    try:
        # data_only: 数式のセルは、最後に Excel で保存したときの計算結果を読む。
        # read_only は使わない。範囲をファイル内の dimension 記録に頼るため、
        # 記録が古いファイル（他のツールで作ったものなど）では行を取りこぼす。
        wb = load_workbook(path, data_only=True)
    except Exception as e:  # 壊れたファイル・別形式など、openpyxl の例外は多岐にわたる
        raise SheetError(f"Excel を開けません: {path.name}（{e}）") from e

    try:
        sheetnames = list(wb.sheetnames)
        result = Book()
        for name in (REGISTER_SHEET, UPDATE_SHEET):
            if name in sheetnames:
                result[name] = _read_sheet(wb[name], name)
        if PROJECT_SHEET in sheetnames:
            result.project_key = _read_project_key(wb[PROJECT_SHEET])
    finally:
        wb.close()

    if not result:
        raise SheetError(
            f"「{REGISTER_SHEET}」「{UPDATE_SHEET}」のどちらのシートもありません: {path.name}"
            f"（シート名: {' / '.join(sheetnames)}）"
        )
    return result


def _read_project_key(ws) -> str:
    """A 列が「プロジェクトキー」の行の、B 列の値。その行が無ければ ""。"""
    for row in ws.iter_rows(min_col=1, max_col=2, values_only=True):
        label, value = (tuple(row) + (None, None))[:2]
        if not is_blank(label) and str(label).strip() == PROJECT_LABEL:
            return "" if is_blank(value) else str(value).strip()
    return ""


def _read_sheet(ws, name: str) -> SheetData:
    rows_iter = ws.iter_rows(values_only=True)
    header_row = next(rows_iter, None) or ()
    headers = []
    for idx, value in enumerate(header_row, start=1):
        if is_blank(value):
            continue
        header = str(value).strip()
        if header.startswith(COMMENT_PREFIX):
            continue
        headers.append((idx, header))

    data = SheetData(name=name, headers=headers)
    wanted = [idx for idx, _ in headers]
    for row_no, row in enumerate(rows_iter, start=2):
        values = {idx: (row[idx - 1] if idx - 1 < len(row) else None) for idx in wanted}
        if all(is_blank(v) for v in values.values()):
            continue
        data.rows.append(SheetRow(row_no=row_no, values=values))
    return data


# ---------------------------------------------------------------------------
# 書き込み
# ---------------------------------------------------------------------------

HEADER_FONT = Font(bold=True, color="FFFFFF")
HEADER_FILL = PatternFill("solid", fgColor="4F6D8F")
STRUCTURE_FILL = PatternFill("solid", fgColor="8F5F4F")   # 課題キー・親子の列
REFERENCE_FILL = PatternFill("solid", fgColor="8A8A8A")   # # で始まる参照用の列
WRAP = Alignment(wrap_text=True, vertical="top")
TOP = Alignment(vertical="top")

COLUMN_WIDTHS = {
    KEY_HEADER: 12, PARENT_HEADER: 12, CHILD_HEADER: 8,
    "件名": 40, "詳細": 50, "#説明": 50,
}


def register_headers(master: Master) -> list[str]:
    return [PARENT_HEADER, CHILD_HEADER] + [c.header for c in all_columns(master)]


def update_headers(master: Master) -> list[str]:
    return [KEY_HEADER, PARENT_HEADER] + [c.header for c in all_columns(master)]


def _style_header(ws, headers: list[str]) -> None:
    for idx, header in enumerate(headers, start=1):
        cell = ws.cell(row=1, column=idx, value=header)
        cell.font = HEADER_FONT
        if header.startswith(COMMENT_PREFIX):
            cell.fill = REFERENCE_FILL
        elif header in (KEY_HEADER, PARENT_HEADER, CHILD_HEADER):
            cell.fill = STRUCTURE_FILL
        else:
            cell.fill = HEADER_FILL
        width = COLUMN_WIDTHS.get(header, max(10, min(30, len(header) * 2 + 4)))
        ws.column_dimensions[get_column_letter(idx)].width = width
    ws.freeze_panes = "A2"


def _write_master_sheet(wb: Workbook, master: Master) -> dict[str, str]:
    """
    プルダウンの選択肢を隠しシートに並べる。

    返り値は {見出し: 参照式}。入力規則に直接選択肢を書くと 255 文字の
    上限があり、ユーザーやカテゴリーが多いと収まらないため、範囲で参照する。
    """
    ws = wb.create_sheet(MASTER_SHEET)
    lists: dict[str, list[str]] = {CHILD_HEADER: [CHILD_MARK]}
    for col in all_columns(master):
        if col.kind != "name":   # 複数値の列は 1 つしか選べないプルダウンに向かない
            continue
        table = col.custom.items if col.is_custom else master.table(col.table)
        lists[col.header] = table.names()

    refs = {}
    for idx, (header, names) in enumerate(lists.items(), start=1):
        letter = get_column_letter(idx)
        ws.cell(row=1, column=idx, value=header)
        for r, name in enumerate(names, start=2):
            ws.cell(row=r, column=idx, value=name)
        if names:
            refs[header] = f"'{MASTER_SHEET}'!${letter}$2:${letter}${len(names) + 1}"
    ws.sheet_state = "hidden"
    return refs


def _add_validations(ws, headers: list[str], refs: dict[str, str]) -> None:
    """
    プルダウンを付ける。入力の手助けであって制限ではないため、
    一覧に無い値もエラーにしない（(削除) を書けるように。名前の検査は取り込み時に行う）。
    """
    for idx, header in enumerate(headers, start=1):
        if header not in refs:
            continue
        letter = get_column_letter(idx)
        dv = DataValidation(type="list", formula1=refs[header], allow_blank=True)
        dv.showErrorMessage = False
        dv.add(f"{letter}2:{letter}{VALIDATION_ROWS + 1}")
        ws.add_data_validation(dv)


def _format_cells(ws, headers: list[str], columns: dict[str, Column], n_rows: int) -> None:
    for idx, header in enumerate(headers, start=1):
        col = columns.get(header)
        wrap = col is not None and (col.kind == "names" or header == "詳細")
        is_date = col is not None and col.kind == "date"
        for r in range(2, n_rows + 2):
            cell = ws.cell(row=r, column=idx)
            cell.alignment = WRAP if wrap else TOP
            if is_date:
                cell.number_format = "yyyy-mm-dd"


GUIDE_LINES = [
    "このブックは backlog-issue-sheet で取り込める書式です。",
    "",
    f"■ 「{REGISTER_SHEET}」シート：課題を新しく作る",
    f"  ・{CHILD_HEADER}に {CHILD_MARK} を付けた行は、上へさかのぼって最初に見つかった"
    f" {CHILD_MARK} の無い行の子課題になります。",
    f"  ・既にある課題の子にするときは、{PARENT_HEADER}にその課題キーを書きます。",
    "  ・件名・種別・優先度は必須です。",
    "",
    f"■ 「{UPDATE_SHEET}」シート：既存の課題を変更する",
    f"  ・{KEY_HEADER}は必須です。",
    "  ・空のセルは変更しません。値を消すときは "
    f"{CLEAR_TOKEN} と書きます。",
    f"  ・{PARENT_HEADER}に課題キーを書くとその子課題に、{DETACH_TOKEN} と書くと"
    "子課題ではなくなります。",
    "",
    f"■ 「{PROJECT_SHEET}」シート：どのプロジェクト向けのブックか",
    f"  ・{PROJECT_LABEL}を書きます。取り込むとき、設定ファイルの project_key と違えば何も送りません。",
    "  ・backlog-issue-sheet template で作ったひな形と、書き出したファイルには、はじめから入っています。",
    "",
    "■ 共通",
    "  ・列は見出しの名前で読みます。要らない列は消して構いません。",
    "  ・見出しが # で始まる列は読み飛ばします（メモ用）。",
    "  ・カテゴリーなど複数の値は、改行（Alt+Enter）か「,」で区切ります。",
    "  ・カスタム属性は、その名前を見出しにした列を足すと使えます。",
    "  ・使える名前（種別・担当者・カスタム属性など）は backlog-issue-sheet master で確かめられます。",
    f"  ・{REGISTER_SHEET}シートでは行の順番が親子関係を表すため、並べ替えないでください。",
]


EXAMPLE_NOTE_HEADER = "#説明"


def _write_project_sheet(wb: Workbook, master: Master) -> None:
    ws = wb.create_sheet(PROJECT_SHEET)
    label = ws.cell(row=1, column=1, value=PROJECT_LABEL)
    label.font = HEADER_FONT
    label.fill = STRUCTURE_FILL
    ws.cell(row=1, column=2, value=master.project_key or None)
    ws.cell(row=3, column=1, value=PROJECT_NOTE)
    ws.column_dimensions["A"].width = 16
    ws.column_dimensions["B"].width = 16


def write_workbook(path: str | Path, master: Master, **kwargs) -> Path:
    """ブックを組み立てて保存する。引数は build_workbook と同じ。"""
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    build_workbook(master, **kwargs).save(path)
    return path


def build_workbook(
    master: Master,
    *,
    register_rows: list[dict] | None = None,
    update_rows: list[dict] | None = None,
    extra_update_headers: list[str] | None = None,
    include_register: bool = True,
    include_guide: bool = False,
    examples: dict[str, tuple[str, list[dict]]] | None = None,
) -> Workbook:
    """
    ブックを組み立てる。行は {見出し: 値} の dict で渡す。

    include_register=False なら「更新」シートだけを出す（書き出し・実行結果用）。
    examples は {シート名: (書式にするシート名, 行)}。記入例のシートで、
    「登録」「更新」以外の名前なので取り込みでは読まれない。
    """
    columns = {c.header: c for c in all_columns(master)}

    wb = Workbook()
    wb.remove(wb.active)
    if include_guide:
        guide = wb.create_sheet(GUIDE_SHEET)
        for r, line in enumerate(GUIDE_LINES, start=1):
            guide.cell(row=r, column=1, value=line)
        guide.column_dimensions["A"].width = 100

    sheets = []
    if include_register:
        sheets.append((wb.create_sheet(REGISTER_SHEET), register_headers(master), register_rows or []))
    sheets.append((
        wb.create_sheet(UPDATE_SHEET),
        update_headers(master) + list(extra_update_headers or []),
        update_rows or [],
    ))
    _write_project_sheet(wb, master)
    for name, (based_on, rows) in (examples or {}).items():
        base = register_headers(master) if based_on == REGISTER_SHEET else update_headers(master)
        sheets.append((wb.create_sheet(name), base + [EXAMPLE_NOTE_HEADER], rows))

    refs = _write_master_sheet(wb, master)
    for ws, headers, rows in sheets:
        _style_header(ws, headers)
        for r, row in enumerate(rows, start=2):
            for idx, header in enumerate(headers, start=1):
                value = row.get(header)
                if value is not None:
                    ws.cell(row=r, column=idx, value=value)
        _format_cells(ws, headers, columns, max(len(rows), 1))
        _add_validations(ws, headers, refs)

    wb.active = wb.sheetnames.index(sheets[0][0].title)
    return wb


def timestamp(now: dt.datetime | None = None) -> str:
    return (now or dt.datetime.now()).strftime("%Y%m%d_%H%M%S")


def unique_stamp(directory: Path, names: list[str], now: dt.datetime | None = None) -> str:
    """
    ファイル名に使う日時。names（"run_{}.csv" のような形）のどれかが既にあれば、
    _2, _3 … を付けて、どれとも重ならないものにする。

    日時は秒までなので、同じ秒に続けて実行すると同じ名前になる。重なったまま
    保存すると、前のファイル（どこまで反映したかを示す実行ログなど）が黙って消える。
    実行ログと結果のように対になるファイルは、同じ番号にそろえる。
    """
    base = timestamp(now)
    stamp, n = base, 2
    while any((Path(directory) / name.format(stamp)).exists() for name in names):
        stamp, n = f"{base}_{n}", n + 1
    return stamp
