"""
どのプロジェクト向けのブックか（「プロジェクト」シート）
=====================================================
課題キーを書かない「登録」シートは、設定の接続先が別のプロジェクトを指していても
名前が通ってしまい、そのまま作成される。ブックに書いたプロジェクトキーで止める。
"""

from openpyxl import load_workbook

from exporter import ExportFilter, export
from planner import build_plan
from sheet_io import PROJECT_LABEL, PROJECT_SHEET, read_workbook, write_workbook

REG = {"登録": (["件名", "種別", "優先度"], [["A", "タスク", "中"]])}


def book_for(project_key):
    """登録シートに、プロジェクトシート（A1 見出し・B1 値）を足す。"""
    return {**REG, PROJECT_SHEET: ([PROJECT_LABEL, project_key], [])}


def project_messages(items):
    return [str(p) for p in items if p.sheet == PROJECT_SHEET]


class TestCheck:
    def test_一致すれば注意もエラーも出ない(self, plan_of):
        plan = plan_of(book_for("DEMO"))
        assert plan.ok
        assert project_messages(plan.warnings) == []

    def test_大文字小文字の違いは同じとみなす(self, plan_of):
        assert plan_of(book_for(" demo ")).ok

    def test_違えばエラーにして何も作らない(self, plan_of):
        plan = plan_of(book_for("OTHER"))
        assert not plan.ok
        assert project_messages(plan.problems) == [
            "プロジェクト「プロジェクトキー」: このブックは OTHER 向けですが、設定の接続先は DEMO です。"
            "設定ファイルか、ブックの指定を確かめてください"
        ]

    def test_シートが無ければ注意にとどめる(self, plan_of):
        plan = plan_of(REG)
        assert plan.ok
        assert len(project_messages(plan.warnings)) == 1
        assert "シートが無い" in project_messages(plan.warnings)[0]

    def test_値が空なら注意にとどめる(self, plan_of):
        plan = plan_of(book_for(None))
        assert plan.ok
        assert "空のため" in project_messages(plan.warnings)[0]

    def test_見出しの行が無ければ空とみなす(self, plan_of):
        plan = plan_of({**REG, PROJECT_SHEET: (["メモ", "DEMO"], [])})
        assert plan.ok
        assert "空のため" in project_messages(plan.warnings)[0]


class TestWrite:
    def test_ひな形にはプロジェクトキーが入る(self, master, tmp_path):
        path = write_workbook(tmp_path / "t.xlsx", master, include_guide=True)
        ws = load_workbook(path)[PROJECT_SHEET]
        assert (ws["A1"].value, ws["B1"].value) == (PROJECT_LABEL, "DEMO")
        assert read_workbook(path).project_key == "DEMO"

    def test_書き出したファイルを戻しても注意が出ない(self, fake, master, tmp_path):
        fake.add("x")
        out, _, _ = export(fake, master, tmp_path / "o.xlsx", ExportFilter())
        plan = build_plan(read_workbook(out), master, fake)
        assert plan.ok and project_messages(plan.warnings) == []

