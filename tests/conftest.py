"""
テスト共通のフィクスチャ
========================
FakeBacklog は課題の状態を保持し、Backlog の親子の制約（1 階層まで）もまねる。
実行の順番を誤ると、本物と同じように更新が拒否される。
"""

from __future__ import annotations

import pytest
from openpyxl import Workbook

import backlog_client
from backlog_client import BacklogAPIError
from master import Master

PROJECT = {"id": 42, "projectKey": "DEMO", "subtaskingEnabled": True}
ISSUE_TYPES = [{"id": 1, "name": "タスク"}, {"id": 2, "name": "バグ"}]
PRIORITIES = [{"id": 2, "name": "高"}, {"id": 3, "name": "中"}, {"id": 4, "name": "低"}]
STATUSES = [
    {"id": 1, "name": "未対応"}, {"id": 2, "name": "処理中"},
    {"id": 3, "name": "処理済み"}, {"id": 4, "name": "完了"},
]
RESOLUTIONS = [{"id": 0, "name": "対応済み"}, {"id": 1, "name": "対応しない"}]
USERS = [
    {"id": 10, "name": "山田太郎", "userId": "yamada"},
    {"id": 11, "name": "佐藤花子", "userId": "sato"},
]
CATEGORIES = [{"id": 100, "name": "設計"}, {"id": 101, "name": "開発"}]
VERSIONS = [{"id": 200, "name": "v1.0"}, {"id": 201, "name": "v2.0"}]
CUSTOM_FIELDS = [
    {"id": 300, "typeId": 1, "name": "顧客名"},
    {"id": 301, "typeId": 3, "name": "見積"},
    {"id": 302, "typeId": 5, "name": "区分", "items": [{"id": 1, "name": "A"}, {"id": 2, "name": "B"}]},
    {"id": 303, "typeId": 6, "name": "タグ", "items": [{"id": 11, "name": "x"}, {"id": 12, "name": "y"}]},
    {"id": 304, "typeId": 4, "name": "締切"},
    {"id": 305, "typeId": 1, "name": "バグ専用", "applicableIssueTypes": [2]},
]

_LIST_FIELDS = {"categoryId": ("category", CATEGORIES), "versionId": ("versions", VERSIONS),
                "milestoneId": ("milestone", VERSIONS)}
_SINGLE_FIELDS = {"issueTypeId": ("issueType", ISSUE_TYPES), "priorityId": ("priority", PRIORITIES),
                  "statusId": ("status", STATUSES), "resolutionId": ("resolution", RESOLUTIONS),
                  "assigneeId": ("assignee", USERS)}


def _obj(items, item_id):
    for item in items:
        if item["id"] == int(item_id):
            return dict(item)
    raise BacklogAPIError(f"不正な ID: {item_id}", status=400)


class FakeBacklog:
    base_url = "https://example.backlog.com/api/v2"

    def __init__(self, *, subtasking=True):
        self.project = dict(PROJECT, subtaskingEnabled=subtasking)
        self.issues: dict[int, dict] = {}
        self.calls: list[tuple] = []         # ("create", params) / ("update", key, params)
        self.fail_create_matching: str | None = None
        self.fatal_on_create = False
        self.no_change_keys: set[str] = set()
        self._next_id = 1000
        self._next_no = 1

    # ---- マスタ ----
    def get_project(self, key): return dict(self.project)
    def get_issue_types(self, key): return ISSUE_TYPES
    def get_priorities(self): return PRIORITIES
    def get_statuses(self, key): return STATUSES
    def get_resolutions(self): return RESOLUTIONS
    def get_project_users(self, key): return USERS
    def get_categories(self, key): return CATEGORIES
    def get_versions(self, key): return VERSIONS
    def get_custom_fields(self, key): return CUSTOM_FIELDS

    # ---- 課題 ----
    def add(self, summary, *, parent=None, **params) -> dict:
        """テスト用に既存の課題を作る（呼び出しは記録しない）。"""
        params.setdefault("issueTypeId", 1)
        params.setdefault("priorityId", 3)
        if parent is not None:
            params["parentIssueId"] = parent["id"]
        issue = self._new_issue(summary)
        self._apply(issue, params)
        return issue

    def _new_issue(self, summary) -> dict:
        self._next_id += 1
        issue = {
            "id": self._next_id, "issueKey": f"DEMO-{self._next_no}", "summary": summary,
            "description": "", "issueType": None, "priority": None, "status": {"id": 1, "name": "未対応"},
            "resolution": None, "assignee": None, "category": [], "versions": [], "milestone": [],
            "startDate": None, "dueDate": None, "estimatedHours": None, "actualHours": None,
            "parentIssueId": None, "customFields": [], "updated": "2026-10-01T09:00:00Z",
        }
        self._next_no += 1
        self.issues[issue["id"]] = issue
        return issue

    def _find(self, key_or_id) -> dict | None:
        text = str(key_or_id)
        for issue in self.issues.values():
            if issue["issueKey"] == text or str(issue["id"]) == text:
                return issue
        return None

    def _children_of(self, issue_id):
        return [i for i in self.issues.values() if i["parentIssueId"] == issue_id]

    def _apply(self, issue: dict, params: dict) -> None:
        for key, value in params.items():
            if key in ("projectId",):
                continue
            if key == "parentIssueId":
                if value in ("", None):
                    issue["parentIssueId"] = None
                    continue
                parent = self._find(value)
                if parent is None:
                    raise BacklogAPIError("親課題が見つかりません", status=400)
                if parent["parentIssueId"]:
                    raise BacklogAPIError("子課題は親にできません", status=400)
                if self._children_of(issue["id"]):
                    raise BacklogAPIError("子課題を持つ課題は子にできません", status=400)
                issue["parentIssueId"] = parent["id"]
            elif key in _SINGLE_FIELDS:
                attr, items = _SINGLE_FIELDS[key]
                issue[attr] = None if value == "" else _obj(items, value)
            elif key in _LIST_FIELDS:
                attr, items = _LIST_FIELDS[key]
                issue[attr] = [_obj(items, v) for v in value if v != ""]
            elif key in ("startDate", "dueDate"):
                issue[key] = f"{value}T00:00:00Z" if value else None
            elif key in ("estimatedHours", "actualHours"):
                issue[key] = float(value) if value != "" else None
            elif key in ("summary", "description"):
                issue[key] = value
            elif key.startswith("customField_"):
                self._apply_custom(issue, int(key.removeprefix("customField_")), value)
            else:
                raise AssertionError(f"未知のパラメータ: {key}")

    def _apply_custom(self, issue, cf_id, value):
        cf = next(c for c in CUSTOM_FIELDS if c["id"] == cf_id)
        issue["customFields"] = [c for c in issue["customFields"] if c["id"] != cf_id]
        if value in ("", [""], []):
            stored = None
        elif cf["typeId"] in (5, 8):
            stored = _obj(cf["items"], value)
        elif cf["typeId"] in (6, 7):
            stored = [_obj(cf["items"], v) for v in value]
        elif cf["typeId"] == 3:
            stored = float(value)
        else:
            stored = value
        issue["customFields"].append({"id": cf_id, "fieldTypeId": cf["typeId"], "name": cf["name"], "value": stored})

    def get_issue(self, key_or_id):
        issue = self._find(key_or_id)
        return dict(issue) if issue else None

    def get_issues(self, project_id, params=None):
        params = params or {}
        result = []
        for issue in self.issues.values():
            if "statusId" in params and issue["status"]["id"] not in params["statusId"]:
                continue
            if "issueTypeId" in params and issue["issueType"]["id"] not in params["issueTypeId"]:
                continue
            if "parentIssueId" in params and issue["parentIssueId"] not in params["parentIssueId"]:
                continue
            if "keyword" in params and params["keyword"] not in issue["summary"]:
                continue
            result.append(dict(issue))
        return result

    def get_child_issues(self, project_id, parent_issue_id):
        return [dict(i) for i in self._children_of(parent_issue_id)]

    def create_issue(self, params):
        self.calls.append(("create", dict(params)))
        if self.fatal_on_create:
            raise BacklogAPIError("認証エラー", status=401, fatal=True)
        if self.fail_create_matching and self.fail_create_matching in params["summary"]:
            raise BacklogAPIError("作成に失敗しました", status=400)
        issue = self._new_issue(params["summary"])
        try:
            self._apply(issue, params)
        except BacklogAPIError:
            del self.issues[issue["id"]]
            raise
        return dict(issue)

    def update_issue(self, key, params):
        self.calls.append(("update", key, dict(params)))
        if key in self.no_change_keys:
            raise backlog_client.BacklogNoChangeError("変更されていません")
        issue = self._find(key)
        if issue is None:
            raise BacklogAPIError("課題がありません", status=404)
        self._apply(issue, params)
        return dict(issue)


@pytest.fixture(autouse=True)
def no_sleep(monkeypatch):
    monkeypatch.setattr(backlog_client.time, "sleep", lambda _: None)


@pytest.fixture
def fake():
    return FakeBacklog()


@pytest.fixture
def master(fake):
    return Master.load(fake, "DEMO")


@pytest.fixture
def make_book(tmp_path):
    """
    シートごとの (見出し, 行) から xlsx を作る。

        path = make_book({"登録": (["件名", "種別", "優先度"], [["A", "タスク", "中"]])})
    """
    counter = {"n": 0}

    def _make(sheets: dict, name: str | None = None):
        counter["n"] += 1
        wb = Workbook()
        wb.remove(wb.active)
        for sheet_name, (headers, rows) in sheets.items():
            ws = wb.create_sheet(sheet_name)
            for c, h in enumerate(headers, start=1):
                ws.cell(row=1, column=c, value=h)
            for r, row in enumerate(rows, start=2):
                for c, v in enumerate(row, start=1):
                    ws.cell(row=r, column=c, value=v)
        path = tmp_path / (name or f"book{counter['n']}.xlsx")
        wb.save(path)
        return path

    return _make


@pytest.fixture
def plan_of(make_book, master, fake):
    """シートの内容から計画を作る。"""
    from planner import build_plan
    from sheet_io import read_workbook

    def _plan(sheets: dict):
        return build_plan(read_workbook(make_book(sheets)), master, fake)

    return _plan
