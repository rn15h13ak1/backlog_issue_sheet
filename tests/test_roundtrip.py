"""
書き出し → 取り込みの往復
=========================
書き出した Excel を何も変えずに取り込みに戻したとき、差分が 1 件も出ないこと。
ここが崩れると、書き出して一部だけ直したつもりが、ほかの項目まで書き換わる。
"""

import datetime as dt

from openpyxl import load_workbook

from exporter import ExportFilter, export
from planner import build_plan
from sheet_io import read_workbook, write_workbook


def full_issue(fake):
    return fake.add(
        "件名, カンマ入り", description="1行目\r\n2行目", issueTypeId=2, priorityId=2,
        statusId=3, resolutionId=0, assigneeId=11, categoryId=[100, 101], versionId=[200],
        milestoneId=[201], startDate="2026-10-01", dueDate="2026-10-31",
        estimatedHours="1.5", actualHours="2",
        customField_300="ACME", customField_301="3.25", customField_302=2,
        customField_303=[11, 12], customField_304="2026-12-01", customField_305="バグの値",
    )


class TestRoundTrip:
    def test_書き出してそのまま戻すと差分が出ない(self, fake, master, tmp_path):
        p = full_issue(fake)
        fake.add("子", parent=p)
        fake.add("空の課題")
        out, count, _ = export(fake, master, tmp_path / "out.xlsx", ExportFilter())
        assert count == 3

        plan = build_plan(read_workbook(out), master, fake)
        assert plan.ok, [str(x) for x in plan.problems]
        assert [u.issue_key for u in plan.updates if u.has_changes] == []

    def test_一部を直すとそこだけが差分になる(self, fake, master, tmp_path):
        full_issue(fake)
        out, _, _ = export(fake, master, tmp_path / "out.xlsx", ExportFilter())
        wb = load_workbook(out)
        ws = wb["更新"]
        headers = [c.value for c in ws[1]]
        ws.cell(row=2, column=headers.index("期限日") + 1, value=dt.date(2026, 11, 30))
        wb.save(out)

        plan = build_plan(read_workbook(out), master, fake)
        changes = [(c.column.header, c.new) for c in plan.updates[0].changes]
        assert changes == [("期限日", "2026-11-30")]


class TestExport:
    def test_親の直後に子が並び_親課題キーが入る(self, fake, master, tmp_path):
        a = fake.add("A")       # DEMO-1
        fake.add("B")           # DEMO-2
        fake.add("A-1", parent=a)  # DEMO-3
        out, _, _ = export(fake, master, tmp_path / "o.xlsx", ExportFilter())
        ws = load_workbook(out)["更新"]
        rows = [(r[0].value, r[1].value) for r in ws.iter_rows(min_row=2)]
        assert rows == [("DEMO-1", None), ("DEMO-3", "DEMO-1"), ("DEMO-2", None)]

    def test_既定では完了を除く(self, fake, master, tmp_path):
        fake.add("open")
        fake.add("closed", statusId=4)
        _, count, _ = export(fake, master, tmp_path / "o.xlsx", ExportFilter())
        assert count == 1
        _, count, _ = export(fake, master, tmp_path / "o.xlsx", ExportFilter(include_closed=True))
        assert count == 2

    def test_子だけが当たっても親を加える(self, fake, master, tmp_path):
        a = fake.add("親")
        fake.add("対象の子", parent=a)
        _, count, _ = export(fake, master, tmp_path / "o.xlsx", ExportFilter(keyword="対象"))
        assert count == 2

    def test_親を指定して書き出す(self, fake, master, tmp_path):
        a = fake.add("親")
        fake.add("子", parent=a)
        fake.add("無関係")
        _, count, _ = export(fake, master, tmp_path / "o.xlsx", ExportFilter(parent_key="DEMO-1"))
        assert count == 2

    def test_長すぎる本文は空にする(self, fake, master, tmp_path):
        fake.add("x", description="あ" * 40000)
        out, _, warnings = export(fake, master, tmp_path / "o.xlsx", ExportFilter())
        assert any("上限" in w for w in warnings)
        # 空のセルは「変更しない」なので、戻しても本文は消えない
        plan = build_plan(read_workbook(out), master, fake)
        assert not plan.updates[0].has_changes

    def test_参照用の列は取り込みで無視される(self, fake, master, tmp_path):
        fake.add("x")
        out, _, _ = export(fake, master, tmp_path / "o.xlsx", ExportFilter())
        ws = load_workbook(out)["更新"]
        headers = [c.value for c in ws[1]]
        assert "#URL" in headers
        url = ws.cell(row=2, column=headers.index("#URL") + 1).value
        assert url == "https://example.backlog.com/view/DEMO-1"


class TestTemplate:
    def test_ひな形は空のまま取り込める(self, master, fake, tmp_path):
        path = write_workbook(tmp_path / "t.xlsx", master, include_guide=True)
        wb = load_workbook(path)
        assert wb.sheetnames[:3] == ["使い方", "登録", "更新"]
        assert wb["_マスタ"].sheet_state == "hidden"
        reg = [c.value for c in wb["登録"][1]]
        assert reg[:3] == ["親課題キー", "子課題", "件名"]
        assert "顧客名" in reg and "課題キー" not in reg
        assert [c.value for c in wb["更新"][1]][:2] == ["課題キー", "親課題キー"]

        plan = build_plan(read_workbook(path), master, fake)
        assert plan.ok and not plan.creates and not plan.updates

    def test_プルダウンが付く(self, master, tmp_path):
        path = write_workbook(tmp_path / "t.xlsx", master)
        ws = load_workbook(path)["登録"]
        formulas = [dv.formula1 for dv in ws.data_validations.dataValidation]
        assert any("_マスタ" in f for f in formulas)

    def test_同姓同名の担当者はプルダウンにログイン_ID_で並ぶ(self, fake, tmp_path, monkeypatch):
        """同じ名前を選ぶと、取り込みで「名前では特定できません」になる。選んで通る値だけを並べる。"""
        import conftest
        from master import Master

        monkeypatch.setattr(conftest, "USERS", conftest.USERS + [{"id": 12, "name": "山田太郎", "userId": "yamada2"}])
        master = Master.load(fake, "DEMO")
        path = write_workbook(tmp_path / "t.xlsx", master)

        wb = load_workbook(path)
        ws = wb["_マスタ"]
        column = [c.value for c in ws[1]].index("担当者") + 1
        choices = [ws.cell(row=r, column=column).value for r in range(2, ws.max_row + 1)]
        choices = [c for c in choices if c]
        assert choices == ["yamada", "佐藤花子", "yamada2"]

        # 選択肢をそのまま書いたら、それぞれ別の担当者として取り込める
        reg = wb["登録"]
        headers = [c.value for c in reg[1]]
        for r, choice in enumerate(choices, start=2):
            for header, value in (("件名", choice), ("種別", "タスク"), ("優先度", "中"), ("担当者", choice)):
                reg.cell(row=r, column=headers.index(header) + 1, value=value)
        wb.save(path)
        plan = build_plan(read_workbook(path), master, fake)
        assert plan.ok, [str(x) for x in plan.problems]
        assert [c.values["assigneeId"] for c in plan.creates] == [10, 11, 12]
