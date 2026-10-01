"""同梱のひな形（templates/課題シート.xlsx）"""

import sys
from pathlib import Path

import pytest
from openpyxl import Workbook, load_workbook

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT / "scripts"))

import make_default_template as mdt  # noqa: E402
from planner import build_plan  # noqa: E402
from sheet_io import EXAMPLE_NOTE_HEADER, read_workbook  # noqa: E402


def test_同梱のひな形が最新(capsys):
    """生成元を変えたのに作り直し忘れると、古いひな形が配られる。"""
    assert mdt.main(["--check"]) == 0, capsys.readouterr().out


def test_何度作っても同じバイト列になる():
    assert mdt.render() == mdt.render()


def test_空のまま取り込んでもエラーにならない(master, fake):
    plan = build_plan(read_workbook(mdt.OUTPUT), master, fake)
    assert plan.ok and not plan.creates and not plan.updates


def test_記入例のシートは取り込みで読まれない():
    assert list(read_workbook(mdt.OUTPUT)) == ["登録", "更新"]
    assert "記入例（登録）" in load_workbook(mdt.OUTPUT).sheetnames


@pytest.fixture
def examples_as_input(tmp_path, fake):
    """
    記入例を「登録」「更新」シートに写したブック。

    記入例の課題キー（PROJ-…）はテスト用のプロジェクトに合わせて DEMO-… に変え、
    参照される課題を FakeBacklog に用意する。
    """
    fake._next_no = 50
    parent = fake.add("PROJ-50 の代わり")           # DEMO-50
    fake._next_no = 101
    fake.add("101")                                 # DEMO-101
    fake.add("102", parent=parent)                  # DEMO-102
    fake.add("103", assigneeId=10)                  # DEMO-103

    src = load_workbook(mdt.OUTPUT)
    wb = Workbook()
    wb.remove(wb.active)
    for example, target in (("記入例（登録）", "登録"), ("記入例（更新）", "更新")):
        ws = wb.create_sheet(target)
        for row in src[example].iter_rows(values_only=True):
            ws.append([v.replace("PROJ-", "DEMO-") if isinstance(v, str) else v for v in row])
    path = tmp_path / "examples.xlsx"
    wb.save(path)
    return path


def test_記入例はそのまま取り込める(examples_as_input, master, fake):
    plan = build_plan(read_workbook(examples_as_input), master, fake)
    assert plan.ok, [str(p) for p in plan.problems]
    assert [c.parent_row for c in plan.creates] == [None, 2, 2, None, None]
    assert plan.creates[3].parent_key == "DEMO-50"
    assert all(u.has_changes for u in plan.updates)


def test_記入例には説明の列がある():
    ws = load_workbook(mdt.OUTPUT)["記入例（登録）"]
    assert EXAMPLE_NOTE_HEADER in [c.value for c in ws[1]]
